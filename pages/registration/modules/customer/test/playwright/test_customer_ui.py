import os
import random
import tempfile
import pytest
import openpyxl
from pages.registration.modules.customer.data.customer_data import (
    generate_company_name,
    generate_email,
    generate_phone_number,
    generate_pan_number,
    generate_address,
    generate_luhn_gstin,
    generate_account_number,
    generate_ifsc_code,
)


def _make_data():
    return {
        # Step 1
        "ownership_status":         "Proprietorship",
        "company_name":             generate_company_name(),
        "sale_type":                "Export",
        "supply_type":              "Both",
        "transaction_currency":     "INR",
        "email":                    generate_email(),
        "phone_number":             generate_phone_number(),
        "pan_number":               generate_pan_number(),
        "contact_person":           "Contact Person",
        "preferred_payment_method": "Cash",
        "tax_registration_status":  "Registered",
        "gst_registration_type":    "Regular",
        "payment_terms":            "Immediate",
        "mode_of_delivery":         "Air",
        "delivery_terms":           "Spot",
        "courier_terms":            "Paid",
        # Step 2 — only shipping row needed; billing uses Same as Above
        "address1":                 generate_address(),
        # Step 3
        "bank_name":                "HDFC Bank",
        "bank_branch":              "Pune Branch",
        "bank_ifsc":                generate_ifsc_code(),
        "bank_holder":              "Account Holder",
        "bank_account":             generate_account_number(),
        "account_type":             "Saving",
        "bank_proof":               "Cancelled Cheque",
    }


@pytest.mark.smoke
class TestCustomerCreateAndSearch:
    def test_create_search_view(self, customer_page):
        data = _make_data()

        customer_page.create_record(data)

        customer_page.search_customer(data["company_name"])
        customer_page.verify_customer_exists(data["company_name"])

        customer_page.click_view_button(data["company_name"])
        actual = customer_page.page.locator(
            "xpath=//mat-label[contains(.,'Company Name')]/ancestor::mat-form-field//input"
        ).input_value()
        assert actual == data["company_name"], \
            f"View shows '{actual}', expected '{data['company_name']}'"
        customer_page.force_close_popup()


@pytest.mark.validation
class TestCustomerDuplicateFieldValidation:
    def test_duplicate_pan_rejected_on_edit(self, customer_page):
        existing_pan = customer_page.get_row_field_value(0, "PAN Number")

        data = _make_data()
        customer_page.create_record(data)
        customer_page.search_customer(data["company_name"])

        customer_page.click_edit_button(data["company_name"])
        customer_page._fill_text(customer_page.PAN_NUMBER, existing_pan)
        customer_page.click_update()

        customer_page.page.wait_for_selector(".swal2-container", timeout=10000)
        assert customer_page.page.locator(".swal2-title").text_content().strip() == "Validation Failed"

        with customer_page.page.expect_download() as dl:
            customer_page.page.locator(".swal2-confirm").click()
        customer_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=5000)

        download = dl.value
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            errors = {
                sheet.cell_value(i, 1): sheet.cell_value(i, 3)
                for i in range(1, sheet.nrows)
                if sheet.cell_value(i, 1)
            }
        else:
            wb = openpyxl.load_workbook(tmp_path)
            errors = {
                row[1]: row[3]
                for row in wb.active.iter_rows(min_row=2, values_only=True)
                if row[1]
            }

        assert "PAN Number" in errors, f"PAN Number missing from error file: {errors}"
        assert "already exists" in str(errors["PAN Number"]), \
            f"Unexpected PAN error message: {errors['PAN Number']}"

        customer_page._fill_text(customer_page.PAN_NUMBER, data["pan_number"])
        customer_page.close_popup()
        customer_page.page.wait_for_selector("table#excel-table", timeout=5000)


@pytest.mark.validation
class TestCustomerSubmitExistingPan:
    def test_submit_existing_pan(self, customer_page):
        existing_pan = customer_page.get_first_customer_pan()
        data = _make_data()
        data["pan_number"] = existing_pan
        customer_page.open_add_form()
        customer_page.fill_form(data)
        customer_page.submit()
        customer_page.page.wait_for_selector(".swal2-container", timeout=10000)
        assert customer_page.page.locator(".swal2-title").text_content().strip() == "Validation Failed"
        customer_page.page.evaluate("document.querySelector('.swal2-cancel')?.click()")
        customer_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=5000)


@pytest.mark.regression
class TestCustomerHistoryAccumulation:
    def test_history_grows_with_edits(self, customer_page):
        data = _make_data()
        customer_page.create_record(data)
        customer_page.search_customer(data["company_name"])

        assert customer_page.get_history_entry_count(data["company_name"]) == 0
        customer_page.search_customer(data["company_name"])

        customer_page.edit_field_and_update(
            data["company_name"], "Contact Person Name", "Contact One"
        )
        customer_page.search_customer(data["company_name"])
        assert customer_page.get_history_entry_count(data["company_name"]) == 1
        customer_page.search_customer(data["company_name"])

        customer_page.bulk_edit_field(
            data["company_name"],
            "Contact Person Name",
            ["Contact Two", "Contact Three", "Contact Four", "Contact Five"],
        )
        customer_page.search_customer(data["company_name"])

        assert customer_page.get_history_entry_count(data["company_name"]) == 5


@pytest.mark.validation
class TestCustomerRequiredFieldValidation:
    # Step-1 fields only — address cascade fields (State/District/Taluka/Pin Code)
    # only show errors after advancing to step 2, not on step-1 submit.
    REQUIRED_FIELDS = [
        "Phone Number",
        "PAN Number",
    ]

    def test_empty_submit_shows_required_errors(self, customer_page):
        customer_page.navigate_to_page()
        customer_page.open_add_form()
        customer_page.submit()

        for field in self.REQUIRED_FIELDS:
            loc = customer_page.page.locator(
                f"xpath=//mat-label[contains(.,'{field}')]"
                f"/ancestor::mat-form-field//mat-error[contains(.,'This field is required')]"
            ).first
            assert loc.is_visible(), f"Expected required error for '{field}'"

        customer_page.page.wait_for_selector(".swal2-container", timeout=10000)
        customer_page.page.locator(".swal2-confirm").click()
        customer_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=15000)
        customer_page.close_popup()


@pytest.mark.validation
class TestCustomerRequiredAddressRow:
    def test_delete_random_address_row_triggers_error(self, customer_page):
        customer_page.navigate_to_page()

        # View first row to detect address types and their indices
        customer_page.click_row_action(0, "View")
        customer_page._click_next()

        addr_type_sel = "xpath=//mat-label[contains(.,'Address Type')]/ancestor::mat-form-field//mat-select"
        addr_selects = customer_page.page.locator(addr_type_sel)
        count = addr_selects.count()
        assert count > 0, "No address rows found in View"

        addr_rows = {
            i: addr_selects.nth(i).text_content().strip()
            for i in range(count)
        }
        customer_page.force_close_popup()
        customer_page.page.wait_for_selector("table#excel-table", timeout=5000)

        delete_idx = random.choice(list(addr_rows.keys()))
        deleted_type = addr_rows[delete_idx]

        customer_page.click_row_action(0, "Edit")
        customer_page._click_next()

        customer_page.page.locator("td.action-container button").nth(delete_idx).click()
        customer_page.page.wait_for_timeout(500)

        customer_page.click_update()
        customer_page.page.wait_for_selector(".swal2-container", timeout=10000)
        assert customer_page.page.locator(".swal2-title").text_content().strip() == "Validation Failed"

        with customer_page.page.expect_download() as dl:
            customer_page.page.locator(".swal2-confirm").click()
        customer_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=5000)

        download = dl.value
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            errors = {
                str(sheet.cell_value(i, 1)).strip(): str(sheet.cell_value(i, 3)).strip()
                for i in range(1, sheet.nrows)
                if sheet.cell_value(i, 1)
            }
        else:
            wb = openpyxl.load_workbook(tmp_path)
            errors = {
                str(row[1] or "").strip(): str(row[3] or "").strip()
                for row in wb.active.iter_rows(min_row=2, values_only=True)
                if row[1]
            }

        assert "Address Details" in errors, f"Expected 'Address Details' error row, got: {errors}"
        assert deleted_type.lower() in errors["Address Details"].lower(), \
            f"Expected '{deleted_type}' in error message, got: {errors['Address Details']}"

        customer_page.close_popup()
        customer_page.page.wait_for_selector("table#excel-table", timeout=5000)


@pytest.mark.smoke
class TestCustomerExportMaster:
    def test_export_matches_table(self, customer_page):
        customer_page.navigate_to_page()
        customer_page.set_page_size("1000")

        table_data = customer_page.read_table_data([1, 2, 3])
        assert table_data, "Table must have at least one row before export"

        download = customer_page.export_master()
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        excel_data = []
        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            data_start = next(
                (i + 1 for i in range(sheet.nrows)
                 if str(sheet.cell_value(i, 0)).strip() == "Sr No"),
                None,
            )
            if data_start is not None:
                for i in range(data_start, sheet.nrows):
                    company = str(sheet.cell_value(i, 1)).strip()
                    if not company:
                        continue
                    raw_phone = sheet.cell_value(i, 2)
                    phone = str(int(raw_phone)) if isinstance(raw_phone, float) else str(raw_phone).strip()
                    status = str(sheet.cell_value(i, 3)).strip()
                    excel_data.append((company, phone, status))
        else:
            wb = openpyxl.load_workbook(tmp_path)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            data_start = next(
                (i + 1 for i, row in enumerate(rows) if str(row[0] or "").strip() == "Sr No"),
                None,
            )
            if data_start is not None:
                for row in rows[data_start:]:
                    company = str(row[1] or "").strip()
                    if not company:
                        continue
                    phone = str(row[2] or "").strip()
                    status = str(row[3] or "").strip()
                    excel_data.append((company, phone, status))

        assert data_start is not None, "Could not find 'Sr No' header row in exported Excel"
        assert len(excel_data) == len(table_data), (
            f"Row count mismatch: table={len(table_data)}, excel={len(excel_data)}"
        )
        for i, (t, e) in enumerate(zip(table_data, excel_data)):
            assert t[0] == e[0], f"Row {i+1} Company Name mismatch: table={t[0]!r}, excel={e[0]!r}"
            assert t[1] == e[1], f"Row {i+1} Phone Number mismatch: table={t[1]!r}, excel={e[1]!r}"
            assert t[2] == e[2], f"Row {i+1} Status mismatch: table={t[2]!r}, excel={e[2]!r}"


@pytest.mark.regression
class TestCustomerEditAndHistory:
    def test_create_edit_view_history(self, customer_page):
        data = _make_data()

        customer_page.create_record(data)
        customer_page.search_customer(data["company_name"])
        customer_page.verify_customer_exists(data["company_name"])

        history_before = customer_page.get_history_entry_count(data["company_name"])
        assert history_before == 0, f"Expected 0 history rows after create, got {history_before}"

        updated_contact = "Updated Contact Person"
        customer_page.edit_field_and_update(data["company_name"], "Contact Person Name", updated_contact)
        customer_page.search_customer(data["company_name"])

        actual = customer_page.get_view_field_value(data["company_name"], "Contact Person Name")
        assert actual == updated_contact, f"Expected '{updated_contact}', got '{actual}'"

        history_after = customer_page.get_history_entry_count(data["company_name"])
        assert history_after > 0, "Expected history entries after edit, got none"

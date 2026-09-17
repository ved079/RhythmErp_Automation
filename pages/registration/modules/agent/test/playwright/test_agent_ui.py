import os
import random
import tempfile
import pytest
import openpyxl
from pages.registration.modules.agent.data.agent_data import (
    generate_agent_name,
    generate_email,
    generate_phone_number,
    generate_address,
    generate_ifsc_code,
    generate_account_number,
)


def _make_data():
    return {
        "agent_name":               generate_agent_name(),
        "phone_number":             generate_phone_number(),
        "email":                    generate_email(),
        "state":                    "Maharashtra",
        "address":                  generate_address(),
        "payment_terms":            "Immediate",
        "preferred_payment_method": "Cash",
        "bank_name":                "HDFC Bank",
        "bank_branch":              "Pune Branch",
        "bank_ifsc":                generate_ifsc_code(),
        "account_type":             "Current",
        "bank_holder":              "Account Holder",
        "bank_account":             generate_account_number(),
        "bank_proof":               "Passbook",
    }


@pytest.mark.smoke
class TestAgentCreateAndSearch:
    def test_create_search_view(self, agent_page):
        data = _make_data()

        agent_page.create_record(data)

        agent_page.search_agent(data["agent_name"])
        agent_page.verify_agent_exists(data["agent_name"])

        agent_page.click_view_button(data["agent_name"])
        actual = agent_page.page.locator(
            "xpath=//mat-label[contains(.,'Agent Name')]/ancestor::mat-form-field//input"
        ).input_value()
        assert actual == data["agent_name"], \
            f"View shows '{actual}', expected '{data['agent_name']}'"
        agent_page.force_close_popup()


@pytest.mark.regression
class TestAgentEditAndHistory:
    def test_create_edit_view_history(self, agent_page):
        data = _make_data()
        updated_name = generate_agent_name()

        # create
        agent_page.create_record(data)

        # search + verify created
        agent_page.search_agent(data["agent_name"])
        agent_page.verify_agent_exists(data["agent_name"])

        # edit name
        agent_page.click_edit_button(data["agent_name"])
        loc = agent_page.page.locator(agent_page.AGENT_NAME).first
        loc.click(click_count=3)
        loc.fill(updated_name)
        loc.press("Tab")
        agent_page.click_update()
        agent_page.page.wait_for_timeout(1000)

        # search updated name + view confirms value
        agent_page.search_agent(updated_name)
        agent_page.verify_agent_exists(updated_name)
        agent_page.click_view_button(updated_name)
        actual = agent_page.page.locator(
            "xpath=//mat-label[contains(.,'Agent Name')]/ancestor::mat-form-field//input"
        ).input_value()
        assert actual == updated_name, f"Expected '{updated_name}', got '{actual}'"
        agent_page.force_close_popup()

        # history
        history = agent_page.get_history_entry_count(updated_name)
        assert history > 0, "Expected history entries after edit, got none"


@pytest.mark.regression
class TestAgentHistoryAccumulation:
    def test_history_grows_with_edits(self, agent_page):
        data = _make_data()
        updated_name = generate_agent_name()

        agent_page.create_record(data)
        agent_page.search_agent(data["agent_name"])
        assert agent_page.get_history_entry_count(data["agent_name"]) == 0

        # edit name
        agent_page.click_edit_button(data["agent_name"])
        loc = agent_page.page.locator(agent_page.AGENT_NAME).first
        loc.click(click_count=3)
        loc.fill(updated_name)
        loc.press("Tab")
        agent_page.click_update()
        agent_page.page.wait_for_timeout(1000)

        agent_page.search_agent(updated_name)
        assert agent_page.get_history_entry_count(updated_name) == 1


@pytest.mark.validation
class TestAgentDuplicateNameValidation:
    def test_duplicate_name_shows_error_excel(self, agent_page):
        agent_page.navigate_to_page()
        existing_name = agent_page.get_first_agent_name()

        agent_page.open_add_form()
        data = _make_data()
        data["agent_name"] = existing_name
        agent_page.fill_form(data)
        agent_page.submit()

        path = agent_page.download_validation_error_excel()

        import openpyxl, os
        errors = []
        ws = openpyxl.load_workbook(path).active
        for row in ws.iter_rows(values_only=True):
            errors.append(str(row))
        combined = " ".join(errors).lower()
        assert "already exists" in combined or "agent name" in combined, \
            f"Expected duplicate error in excel, got:\n" + "\n".join(errors)


@pytest.mark.validation
class TestAgentRequiredFieldValidation:
    REQUIRED_FIELDS = [
        "Agent Name",
        "Phone Number",
        "State",
        "District",
        "Taluka",
        "Pin Code",
    ]

    def test_empty_submit_shows_required_errors(self, agent_page):
        agent_page.navigate_to_page()
        agent_page.open_add_form()
        agent_page.submit()

        agent_page.page.wait_for_selector(".swal2-container", timeout=10000)

        for field in self.REQUIRED_FIELDS:
            loc = agent_page.page.locator(
                f"xpath=//mat-form-field[contains(@class,'ng-invalid') and .//mat-label[contains(.,'{field}')]]"
            ).first
            assert loc.count() > 0, f"Expected ng-invalid on form field for '{field}'"

        agent_page.page.evaluate("document.querySelector('.swal2-confirm')?.click()")
        agent_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=15000)
        agent_page.close_popup()


@pytest.mark.smoke
class TestAgentExportMaster:
    def test_export_matches_table(self, agent_page):
        agent_page.navigate_to_page()
        agent_page.set_page_size("1000")

        table_data = agent_page.read_table_data([1, 2, 3])
        assert table_data, "Table must have at least one row before export"

        download = agent_page.export_master()
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        excel_data = []
        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            data_start = next(
                (i + 1 for i in range(sheet.nrows)
                 if str(sheet.cell_value(i, 0)).strip() == "Sr No"),
                None,
            )
            if data_start is not None:
                for i in range(data_start, sheet.nrows):
                    col1 = str(sheet.cell_value(i, 1)).strip()
                    if not col1:
                        continue
                    col2 = str(sheet.cell_value(i, 2)).strip()
                    col3 = str(sheet.cell_value(i, 3)).strip()
                    excel_data.append((col1, col2, col3))
        else:
            wb = openpyxl.load_workbook(tmp_path)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            data_start = next(
                (i + 1 for i, row in enumerate(rows) if str(row[0] or "").strip() == "Sr No"),
                None,
            )
            if data_start is not None:
                for row in rows[data_start:]:
                    col1 = str(row[1] or "").strip()
                    if not col1:
                        continue
                    col2 = str(row[2] or "").strip()
                    col3 = str(row[3] or "").strip()
                    excel_data.append((col1, col2, col3))

        assert data_start is not None, "Could not find 'Sr No' header row in exported Excel"
        assert len(excel_data) == len(table_data), (
            f"Row count mismatch: table={len(table_data)}, excel={len(excel_data)}"
        )
        for i, (t, e) in enumerate(zip(table_data, excel_data)):
            assert t[0] == e[0], f"Row {i+1} col1 mismatch: table={t[0]!r}, excel={e[0]!r}"

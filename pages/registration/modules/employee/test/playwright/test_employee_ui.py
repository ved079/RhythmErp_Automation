import os
import tempfile
import pytest
import openpyxl
from pages.registration.modules.employee.data.employee_data import (
    generate_employee_name,
    generate_email,
    generate_phone_number,
)


def _make_data():
    return {
        "employee_name": generate_employee_name(),
        "phone_number":  generate_phone_number(),
        "email":         generate_email(),
    }


@pytest.mark.smoke
class TestEmployeeCreateAndSearch:
    def test_create_search_view(self, employee_page):
        data = _make_data()

        employee_page.create_record(data)

        employee_page.search_employee(data["employee_name"])
        employee_page.verify_employee_exists(data["employee_name"])

        employee_page.click_view_button(data["employee_name"])
        actual = employee_page.page.locator(
            "xpath=//mat-label[contains(.,'Employee Name')]/ancestor::mat-form-field//input"
        ).input_value()
        assert actual == data["employee_name"], \
            f"View shows '{actual}', expected '{data['employee_name']}'"
        employee_page.force_close_popup()


@pytest.mark.regression
class TestEmployeeEditAndHistory:
    def test_create_edit_view_history(self, employee_page):
        data = _make_data()

        employee_page.create_record(data)

        employee_page.search_employee(data["employee_name"])
        employee_page.verify_employee_exists(data["employee_name"])

        employee_page.click_edit_button(data["employee_name"])
        employee_page._select_random_mat_option(employee_page.DESIGNATION)
        employee_page.page.locator(employee_page.DEPARTMENT).wait_for(state="visible", timeout=10000)
        employee_page._select_random_mat_option(employee_page.DEPARTMENT)
        employee_page.click_update()
        employee_page.page.wait_for_timeout(1000)

        employee_page.search_employee(data["employee_name"])
        employee_page.verify_employee_exists(data["employee_name"])

        history = employee_page.get_history_entry_count(data["employee_name"])
        assert history > 0, "Expected history entries after edit, got none"


@pytest.mark.regression
class TestEmployeeHistoryAccumulation:
    def test_history_grows_with_edits(self, employee_page):
        data = _make_data()

        employee_page.create_record(data)
        employee_page.search_employee(data["employee_name"])
        assert employee_page.get_history_entry_count(data["employee_name"]) == 0

        employee_page.click_edit_button(data["employee_name"])
        employee_page._select_random_mat_option(employee_page.DESIGNATION)
        employee_page.page.locator(employee_page.DEPARTMENT).wait_for(state="visible", timeout=10000)
        employee_page._select_random_mat_option(employee_page.DEPARTMENT)
        employee_page.click_update()
        employee_page.page.wait_for_timeout(1000)

        employee_page.search_employee(data["employee_name"])
        assert employee_page.get_history_entry_count(data["employee_name"]) == 1


@pytest.mark.validation
class TestEmployeeDuplicateNameValidation:
    def test_duplicate_name_shows_error_excel(self, employee_page):
        employee_page.navigate_to_page()
        existing_name = employee_page.get_first_employee_name()

        employee_page.open_add_form()
        data = _make_data()
        data["employee_name"] = existing_name
        employee_page.fill_form(data)
        employee_page.submit()

        path = employee_page.download_validation_error_excel()

        errors = []
        ws = openpyxl.load_workbook(path).active
        for row in ws.iter_rows(values_only=True):
            errors.append(str(row))
        combined = " ".join(errors).lower()
        assert "already exists" in combined or "employee name" in combined, \
            f"Expected duplicate error in excel, got:\n" + "\n".join(errors)


@pytest.mark.validation
class TestEmployeeRequiredFieldValidation:
    REQUIRED_FIELDS = [
        "Employee Name",
        "Email",
        "Phone Number",
        "Designation",
        "Department",
    ]

    def test_empty_submit_shows_required_errors(self, employee_page):
        employee_page.navigate_to_page()
        employee_page.open_add_form()
        employee_page.submit()

        employee_page.page.wait_for_selector(".swal2-container", timeout=10000)

        for field in self.REQUIRED_FIELDS:
            loc = employee_page.page.locator(
                f"xpath=//mat-form-field[contains(@class,'ng-invalid') and .//mat-label[contains(.,'{field}')]]"
            ).first
            assert loc.count() > 0, f"Expected ng-invalid on form field for '{field}'"

        employee_page.page.evaluate("document.querySelector('.swal2-confirm')?.click()")
        employee_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=15000)
        employee_page.close_popup()


@pytest.mark.smoke
class TestEmployeeExportMaster:
    def test_export_matches_table(self, employee_page):
        employee_page.navigate_to_page()
        employee_page.set_page_size("1000")

        table_data = employee_page.read_table_data([1, 2, 3])
        assert table_data, "Table must have at least one row before export"

        download = employee_page.export_master()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        wb = openpyxl.load_workbook(tmp_path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        data_start = next(
            (i + 1 for i, row in enumerate(rows) if str(row[0] or "").strip() == "Sr No"),
            None,
        )
        assert data_start is not None, "Could not find 'Sr No' header row in exported Excel"

        excel_data = []
        for row in rows[data_start:]:
            col1 = str(row[1] or "").strip()
            if not col1:
                continue
            col2 = str(row[2] or "").strip()
            col3 = str(row[3] or "").strip()
            excel_data.append((col1, col2, col3))

        assert len(excel_data) == len(table_data), (
            f"Row count mismatch: table={len(table_data)}, excel={len(excel_data)}"
        )
        for i, (t, e) in enumerate(zip(table_data, excel_data)):
            assert t[0] == e[0], f"Row {i+1} col1 mismatch: table={t[0]!r}, excel={e[0]!r}"

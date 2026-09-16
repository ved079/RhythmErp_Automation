import os
import random
import tempfile
import pytest
import openpyxl
from pages.registration.modules.supplier.data.supplier_data import (
    generate_company_name,
    generate_email,
    generate_phone,
    generate_pan,
    generate_address,
    generate_registration_number,
    generate_ifsc,
    generate_account_number,
    generate_contact_person,
)


def _make_data():
    return {
        "ownership_status": "Proprietorship",
        "company_name":     generate_company_name(),
        "po_type":          "Domestic",
        "email":            generate_email(),
        "phone_number":     generate_phone(),
        "default_currency": "INR",
        "pan_number":       generate_pan(),
        "contact_person":   generate_contact_person(),
        "tax_reg_status":   "Registered",
        "gst_reg_type":     "Regular",
        "state":            "Maharashtra",
        "address":          generate_address(),
        "registration_number": generate_registration_number(),
        "bank_name":        "HDFC Bank",
        "bank_branch":      "Pune Branch",
        "bank_ifsc":        generate_ifsc(),
        "account_type":     "Saving",
        "bank_holder":      generate_contact_person(),
        "bank_account":     generate_account_number(),
        "bank_proof":       "Passbook",
    }


@pytest.mark.smoke
class TestSupplierCreateAndSearch:
    def test_create_and_search(self, supplier_page):
        data = _make_data()
        supplier_page.create_record(data)
        supplier_page.search_supplier(data["company_name"])
        supplier_page.verify_supplier_exists(data["company_name"])


@pytest.mark.validation
class TestSupplierDuplicateFieldValidation:
    def test_duplicate_pan_and_phone_rejected_on_edit(self, supplier_page):
        # Read PAN from an existing supplier (first row)
        existing_pan = supplier_page.get_row_field_value(0, "PAN Number")

        # Create a fresh supplier with unique PAN + phone
        data = _make_data()
        supplier_page.create_record(data)
        supplier_page.search_supplier(data["company_name"])

        # Open edit form
        supplier_page.click_edit_button(data["company_name"])

        # Try saving with duplicate PAN → expect validation error
        supplier_page._fill_text(supplier_page.PAN_NUMBER, existing_pan)
        supplier_page.click_update()
        supplier_page.page.wait_for_selector(".swal2-container", timeout=10000)
        assert supplier_page.page.locator(".swal2-title").text_content().strip() == "Validation Failed"

        # Download errors Excel and assert PAN duplicate message
        with supplier_page.page.expect_download() as dl:
            supplier_page.page.locator(".swal2-confirm").click()
        supplier_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=5000)

        download = dl.value
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            errors = {
                sheet.cell_value(i, 1): sheet.cell_value(i, 3)
                for i in range(1, sheet.nrows)
                if sheet.cell_value(i, 1)
            }
        else:
            wb = openpyxl.load_workbook(tmp_path)
            errors = {
                row[1]: row[3]
                for row in wb.active.iter_rows(min_row=2, values_only=True)
                if row[1]
            }

        assert "PAN Number" in errors, f"PAN Number missing from error file: {errors}"
        assert "already exists" in str(errors["PAN Number"]), \
            f"Unexpected PAN error message: {errors['PAN Number']}"

        # Revert PAN and close without saving
        supplier_page._fill_text(supplier_page.PAN_NUMBER, data["pan_number"])

        # Close edit form without saving — test passes
        supplier_page.close_popup()
        supplier_page.page.wait_for_selector("table#excel-table", timeout=5000)


@pytest.mark.validation
class TestSupplierSubmitExistingPan:
    def test_submit_existing_pan(self, supplier_page):
        existing_pan = supplier_page.get_first_supplier_pan()
        data = _make_data()
        data["pan_number"] = existing_pan
        supplier_page.open_add_form()
        supplier_page.fill_form(data)
        supplier_page.submit()
        supplier_page.page.wait_for_selector(".swal2-container", timeout=10000)
        assert supplier_page.page.locator(".swal2-title").text_content().strip() == "Validation Failed"
        supplier_page.page.evaluate("document.querySelector('.swal2-cancel')?.click()")
        supplier_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=5000)


@pytest.mark.regression
class TestSupplierHistoryAccumulation:
    def test_history_grows_with_edits(self, supplier_page):
        data = _make_data()
        supplier_page.create_record(data)
        supplier_page.search_supplier(data["company_name"])

        # No history before any edit
        assert supplier_page.get_history_entry_count(data["company_name"]) == 0
        supplier_page.search_supplier(data["company_name"])

        # First edit — confirm 1 history entry appears
        supplier_page.edit_field_and_update(
            data["company_name"], "Contact Person Name", "Contact One"
        )
        supplier_page.search_supplier(data["company_name"])
        assert supplier_page.get_history_entry_count(data["company_name"]) == 1
        supplier_page.search_supplier(data["company_name"])

        # 4 more back-to-back edits without checking history in between
        supplier_page.bulk_edit_field(
            data["company_name"],
            "Contact Person Name",
            ["Contact Two", "Contact Three", "Contact Four", "Contact Five"],
        )
        supplier_page.search_supplier(data["company_name"])

        # 1 confirmed + 4 batch = 5 total
        assert supplier_page.get_history_entry_count(data["company_name"]) == 5


@pytest.mark.validation
class TestSupplierRequiredFieldValidation:
    # Step-1 fields only — step-2 fields (State/District/Taluka/Pin Code) are
    # not rendered when submitting from step 1 so they can't show inline errors.
    REQUIRED_FIELDS = [
        "Company Name",
        "Phone Number",
        "PAN Number",
        "Tax Registration Status",
    ]

    def test_empty_submit_shows_required_errors(self, supplier_page):
        supplier_page.navigate_to_page()
        supplier_page.open_add_form()
        supplier_page.submit()

        for field in self.REQUIRED_FIELDS:
            loc = supplier_page.page.locator(
                f"xpath=//mat-label[contains(.,'{field}')]"
                f"/ancestor::mat-form-field//mat-error[contains(.,'This field is required')]"
            ).first
            assert loc.is_visible(), f"Expected required error for '{field}'"

        supplier_page.page.wait_for_selector(".swal2-container", timeout=10000)
        supplier_page.page.locator(".swal2-confirm").click()
        supplier_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=15000)
        supplier_page.close_popup()


@pytest.mark.validation
class TestSupplierRequiredAddressRow:
    def test_delete_random_address_row_triggers_error(self, supplier_page):
        supplier_page.navigate_to_page()

        # ── Step 1: View → read all address types and their row indices ───
        supplier_page.click_row_action(0, "View")
        supplier_page._click_next()

        addr_type_sel = "xpath=//mat-label[contains(.,'Address Type')]/ancestor::mat-form-field//mat-select"
        addr_selects = supplier_page.page.locator(addr_type_sel)
        count = addr_selects.count()
        assert count > 0, "No address rows found in View"

        # Build {index: type_label} map
        addr_rows = {
            i: addr_selects.nth(i).text_content().strip()
            for i in range(count)
        }
        supplier_page.force_close_popup()
        supplier_page.page.wait_for_selector("table#excel-table", timeout=5000)

        # Randomly pick one row to delete
        delete_idx = random.choice(list(addr_rows.keys()))
        deleted_type = addr_rows[delete_idx]

        # ── Step 2: Edit → delete the chosen row ─────────────────────────
        supplier_page.click_row_action(0, "Edit")
        supplier_page._click_next()

        supplier_page.page.locator("td.action-container button").nth(delete_idx).click()
        supplier_page.page.wait_for_timeout(500)

        # Update → validation error expected
        supplier_page.click_update()
        supplier_page.page.wait_for_selector(".swal2-container", timeout=10000)
        assert supplier_page.page.locator(".swal2-title").text_content().strip() == "Validation Failed"

        # Download errors Excel
        with supplier_page.page.expect_download() as dl:
            supplier_page.page.locator(".swal2-confirm").click()
        supplier_page.page.wait_for_selector(".swal2-container", state="hidden", timeout=5000)

        download = dl.value
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        # Parse and assert error row
        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            errors = {
                str(sheet.cell_value(i, 1)).strip(): str(sheet.cell_value(i, 3)).strip()
                for i in range(1, sheet.nrows)
                if sheet.cell_value(i, 1)
            }
        else:
            wb = openpyxl.load_workbook(tmp_path)
            errors = {
                str(row[1] or "").strip(): str(row[3] or "").strip()
                for row in wb.active.iter_rows(min_row=2, values_only=True)
                if row[1]
            }

        assert "Address Details" in errors, f"Expected 'Address Details' error row, got: {errors}"
        assert deleted_type.lower() in errors["Address Details"].lower(), \
            f"Expected '{deleted_type}' in error message, got: {errors['Address Details']}"

        # Close without saving to leave the record intact
        supplier_page.close_popup()
        supplier_page.page.wait_for_selector("table#excel-table", timeout=5000)


@pytest.mark.smoke
class TestSupplierExportMaster:
    def test_export_matches_table(self, supplier_page):
        supplier_page.navigate_to_page()
        supplier_page.set_page_size("1000")

        # Snapshot the listing table (Company Name=col1, Phone Number=col2, Status=col3)
        table_data = supplier_page.read_table_data([1, 2, 3])
        assert table_data, "Table must have at least one row before export"

        # Trigger Export Master and download the file
        download = supplier_page.export_master()
        ext = os.path.splitext(download.suggested_filename)[1].lower()
        tmp_path = os.path.join(tempfile.gettempdir(), download.suggested_filename)
        download.save_as(tmp_path)

        # Parse Excel — 4 meta rows (Company/Date/Name/blank) + 1 header row, then data
        # Find the "Sr No" header row dynamically so the offset is format-agnostic
        excel_data = []
        if ext == ".xls":
            import xlrd
            sheet = xlrd.open_workbook(tmp_path).sheet_by_index(0)
            data_start = next(
                (i + 1 for i in range(sheet.nrows)
                 if str(sheet.cell_value(i, 0)).strip() == "Sr No"),
                None,
            )
            if data_start is not None:
                for i in range(data_start, sheet.nrows):
                    company = str(sheet.cell_value(i, 1)).strip()
                    if not company:
                        continue
                    raw_phone = sheet.cell_value(i, 2)
                    phone = str(int(raw_phone)) if isinstance(raw_phone, float) else str(raw_phone).strip()
                    status = str(sheet.cell_value(i, 3)).strip()
                    excel_data.append((company, phone, status))
        else:
            wb = openpyxl.load_workbook(tmp_path)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            data_start = next(
                (i + 1 for i, row in enumerate(rows) if str(row[0] or "").strip() == "Sr No"),
                None,
            )
            if data_start is not None:
                for row in rows[data_start:]:
                    company = str(row[1] or "").strip()
                    if not company:
                        continue
                    phone = str(row[2] or "").strip()
                    status = str(row[3] or "").strip()
                    excel_data.append((company, phone, status))

        assert data_start is not None, "Could not find 'Sr No' header row in exported Excel"
        assert len(excel_data) == len(table_data), (
            f"Row count mismatch: table={len(table_data)}, excel={len(excel_data)}"
        )
        for i, (t, e) in enumerate(zip(table_data, excel_data)):
            assert t[0] == e[0], f"Row {i+1} Company Name mismatch: table={t[0]!r}, excel={e[0]!r}"
            assert t[1] == e[1], f"Row {i+1} Phone Number mismatch: table={t[1]!r}, excel={e[1]!r}"
            assert t[2] == e[2], f"Row {i+1} Status mismatch: table={t[2]!r}, excel={e[2]!r}"


@pytest.mark.regression
class TestSupplierEditAndHistory:
    def test_create_edit_view_history(self, supplier_page):
        data = _make_data()

        # 1. Create
        supplier_page.create_record(data)
        supplier_page.search_supplier(data["company_name"])
        supplier_page.verify_supplier_exists(data["company_name"])

        # 2. History after create — creation does not log a history row
        history_before = supplier_page.get_history_entry_count(data["company_name"])
        assert history_before == 0, f"Expected 0 history rows after create, got {history_before}"

        # 3. Edit contact person
        updated_contact = "Updated Contact Person"
        supplier_page.edit_field_and_update(data["company_name"], "Contact Person Name", updated_contact)
        supplier_page.search_supplier(data["company_name"])

        # 4. View — field must reflect the update
        actual = supplier_page.get_view_field_value(data["company_name"], "Contact Person Name")
        assert actual == updated_contact, f"Expected '{updated_contact}', got '{actual}'"

        # 5. History after edit — must have at least one entry
        history_after = supplier_page.get_history_entry_count(data["company_name"])
        assert history_after > 0, "Expected history entries after edit, got none"
